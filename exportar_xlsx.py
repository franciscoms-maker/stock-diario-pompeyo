#!/usr/bin/env python3
"""Convierte 'Registro Stock Diario.xlsx' (Hoja Registro, Tabla1) a data.json y lo publica.

Uso: python3 exportar_xlsx.py "/ruta/Registro Stock Diario.xlsx" [--sin-publicar]
Requiere: pip install openpyxl. Solo LEE el Excel; nunca lo modifica.
"""
import datetime as dt
import json
import subprocess
import sys
from pathlib import Path

import openpyxl

AQUI = Path(__file__).resolve().parent
COLS = ["fecha", "grupo", "sucursal", "estado", "registrado", "horaLlegada", "puntualidad"]


def fecha(v):
    if isinstance(v, (dt.datetime, dt.date)):
        return v.strftime("%Y-%m-%d")
    s = str(v).strip()
    if len(s) >= 10 and s[2] in "-/" and s[5] in "-/":  # dd-mm-aaaa
        return f"{s[6:10]}-{s[3:5]}-{s[0:2]}"
    return s[:10]


def hora(v):
    if v is None or v == "":
        return ""
    if isinstance(v, (dt.datetime, dt.time)):
        return v.strftime("%H:%M")
    if isinstance(v, (int, float)):  # fracción del día
        m = round((v % 1) * 1440)
        return f"{m // 60 % 24:02d}:{m % 60:02d}"
    return str(v).strip()[:5]


def registrado(v):
    if isinstance(v, dt.datetime):
        return v.strftime("%Y-%m-%d %H:%M")
    return "" if v is None else str(v).strip()


def leer(path):
    wb = openpyxl.load_workbook(path, read_only=False, data_only=True)
    ws = wb["Registro"] if "Registro" in wb.sheetnames else wb.worksheets[0]
    tabla = ws.tables.get("Tabla1") if hasattr(ws, "tables") else None
    rango = tabla.ref if tabla else ws.dimensions
    filas = list(ws[rango])[1:]  # sin encabezado
    out = []
    for f in filas:
        v = [c.value for c in f] + [None] * 7
        if not v[0] or not v[1] or not v[2]:
            continue
        out.append({
            "fecha": fecha(v[0]), "grupo": str(v[1]).strip(), "sucursal": str(v[2]).strip(),
            "estado": str(v[3] or "").strip(), "registrado": registrado(v[4]),
            "horaLlegada": hora(v[5]), "puntualidad": str(v[6] or "").strip(),
        })
    return out


if __name__ == "__main__":
    if len(sys.argv) < 2:
        sys.exit(__doc__)
    datos = leer(sys.argv[1])
    if not datos:
        sys.exit("El Excel no tiene filas; no se publica nada.")
    destino = AQUI / "data.json"
    destino.write_text(json.dumps(datos, ensure_ascii=False), encoding="utf-8")
    print(f"{len(datos)} filas -> {destino}")
    if "--sin-publicar" not in sys.argv:
        subprocess.run([str(AQUI / "publicar.sh"), str(destino)], check=True)
